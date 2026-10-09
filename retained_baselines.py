"""Recover the original published-method results without recomputation.

Only the two published baseline methods are returned. Historical CaPPI scores
remain in the original source and are not substituted for the new CaPPI runs.
This module is used by the bundled document runtime, not by model training.
"""
import hashlib
from pathlib import Path
import re

from docx import Document

SOURCE = Path(__file__).resolve().parent.parent / 'CaPPI-code/SR投稿/CaPPI Manuscript_8.29.docx'
EXPECTED_SHA = '8d8dc91fd2ee80d01acd0acf9b2f95a38e40685f324ba0bceae0f7d72690968e'


def blocks():
    if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != EXPECTED_SHA:
        raise ValueError('Original manuscript changed; recheck baseline provenance')
    original = Document(SOURCE)
    table = [[cell.text for cell in row.cells] for row in original.tables[3].rows]
    if table[0] != ['Model', 'ACC', 'Sen', 'Spec', 'Prec', 'MCC', 'F1-Score']:
        raise ValueError('Historical Table 3 header changed')
    by_name = {row[0]: row for row in table[1:]}
    names = ['PASNVGA', 'SDNN-PPI']
    if not all(name in by_name for name in names):
        raise ValueError('Missing historical published-method row')
    paragraph = original.paragraphs[122].text
    patterns = {
        'PASNVGA': r'PASNVGA achieved an AUC of (\d+\.\d+) and an AUPR of (\d+\.\d+)',
        'SDNN-PPI': r'values of (\d+\.\d+) and (\d+\.\d+) obtained by SDNN-PPI',
    }
    ranking = []
    for name in names:
        match = re.search(patterns[name], paragraph)
        if match is None:
            raise ValueError('Cannot identify original ranking values for ' + name)
        ranking.append([name, *match.groups()])
    return [
        {'kind': 'p', 'text':
         'The original study also evaluated PASNVGA [7] and SDNN-PPI [8]. Their previously reported '
         'results are retained without recalculation in Tables 5A and 5B. The original manuscript '
         'described these as author-run evaluations on the study dataset using official code or '
         'libraries, rather than numerical benchmarks copied from the two method publications. '
         'The present revision preserves that reporting provenance; it does not independently '
         'verify the historical training runs or assign them to the revised partitions.'},
        {'kind': 'table', 'caption':
         'Table 5A. Retained ranking results for the published methods, copied unchanged from '
         'the original study Results section. AUC and AUPR retain their original labels; no '
         'standard deviations were reported for these values. These methods were not rerun '
         'under the revised protocol.',
         'headers': ['Published method', 'Original AUC', 'Original AUPR'],
         'rows': ranking, 'retained_values': True},
        {'kind': 'table', 'caption':
         'Table 5B. Retained classification results copied unchanged from original Table 3. '
         'Values and reported mean ± SD strings are preserved as recorded, without recomputation. '
         'ACC, accuracy; Sen, sensitivity; Spec, specificity; Prec, precision.',
         'headers': table[0], 'rows': [by_name[name] for name in names],
         'widths': [91, 60, 60, 60, 60, 60, 60], 'retained_values': True},
        {'kind': 'p', 'text':
         'These retained results document the earlier published-method comparison. They are '
         'separate from the 50 newly trained neural runs and 40 new feature/topology-control '
         'evaluations. The revised graph construction, partitions, negative pools and training '
         'procedures differ from the historical analysis. Accordingly, the retained values '
         'are not pooled with the new results or used to calculate improvements, significance '
         'or a common ranking of methods. In particular, the legacy AUPR values are not '
         'relabelled as the newly computed average precision without an established estimator.'},
    ]
